"""REFES 2024: valida establecimientos de salud de Palermo sin reemplazar GCBA."""
import argparse, asyncio, io
import httpx
from openpyxl import load_workbook
from scrapers.shared.db_helpers import ensure_source, get_conn, get_or_create_entity, mark_source_synced, upsert_property
from scrapers.shared.normalizer import normalize_name, normalize_value
from scrapers.shared.usig import geocode_address

URL="http://datos.salud.gob.ar/dataset/336cf4d9-447a-44c4-8e34-0ba1fc293d55/resource/5d5710df-cf3f-4d91-b9c5-aecab1e06018/download/establecimientos-asistenciales-asentados-registro-federal-refes-20241227.xlsx"
ORIGIN="https://datos.gob.ar/dataset/salud-listado-establecimientos-salud-asentados-registro-federal-refes"
def clean(v): return str(v or "").strip()
def first(row,*keys):
    for key in keys:
        if clean(row.get(key)): return clean(row[key])
    return ""

async def main():
    p=argparse.ArgumentParser(); p.add_argument("--write",action="store_true"); p.add_argument("--limit",type=int); args=p.parse_args()
    async with httpx.AsyncClient(timeout=120,follow_redirects=True,verify=False) as client:
        data=(await client.get(URL)).content
    ws=load_workbook(io.BytesIO(data),read_only=True,data_only=True).active
    headers=[clean(v).lower().replace(" ","_") for v in next(ws.iter_rows(values_only=True))]
    rows=[dict(zip(headers,values)) for values in ws.iter_rows(values_only=True)]
    postal_prefixes={"1414","1425","1426","1427","1428"}
    candidates=[r for r in rows if first(r,"provincia_nombre").upper() in {"CIUDAD AUTONOMA DE BUENOS AIRES","CABA"} and first(r,"cp").replace("C","")[:4] in postal_prefixes]
    if args.limit: candidates=candidates[:args.limit]
    print(f"REFES: {len(rows)} filas | {len(candidates)} candidatas por CP Palermo | write={args.write}")
    if not args.write:return
    conn=await get_conn()
    try:
        source=await ensure_source(conn,"refes_historical",ORIGIN,4)
        async with httpx.AsyncClient(timeout=30) as geocoder:
         semaphore=asyncio.Semaphore(5)
         async def locate(row):
          address=first(row,"domicilio","direccion")
          try:
           async with semaphore: return await geocode_address(geocoder,address) if address else None
          except httpx.HTTPError: return None
         geocoded=await asyncio.gather(*(locate(row) for row in candidates))
         for row,geo in zip(candidates,geocoded):
            refes=first(row,"establecimiento_id","id_establecimiento")
            entity=await conn.fetchval("""SELECT p.entity_id FROM active_properties p JOIN entities e ON e.id=p.entity_id
                WHERE p.key='refes_id' AND p.value=$1 AND e.is_active AND e.canonical_id IS NULL LIMIT 1""",refes) if refes else None
            name=normalize_name(first(row,"establecimiento_nombre","nombre"))
            if not name:continue
            address=first(row,"domicilio","direccion")
            if not geo: continue
            if not entity: entity=await get_or_create_entity(conn,name,"Facility","health_facility",geo.lat,geo.lng,origin_url=ORIGIN)
            fields={"refes_id":refes,"health_facility_type":first(row,"tipologia_nombre","tipo_establecimiento","tipo"),"jurisdiction":first(row,"provincia_nombre"),"funding_type":first(row,"origen_financiamiento","financiamiento","sector"),"address":address,"historical_status":"Historical Source","historical_data_year":"2024"}
            for key,value in fields.items():
                if value: await upsert_property(conn,str(entity),key,normalize_value(value),"string",source,[ORIGIN],.75)
        await mark_source_synced(conn,source)
    finally:await conn.close()
if __name__=="__main__":asyncio.run(main())
